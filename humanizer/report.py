"""Analysis reports in .json, .docx and .pdf (RULES.md section 5)."""
from __future__ import annotations

import json
from pathlib import Path
from xml.sax.saxutils import escape

import docx
from docx.shared import Pt, RGBColor

from .models import AnalysisResult, Finding

DISCLAIMER = ("이 리포트는 문체 개선을 위한 참고 자료입니다. 원고 작성에 AI를 사용했다면, "
              "투고 저널의 AI 사용 공개 규정을 따로 지켜야 합니다.")
_RISK_RGB = {"high": RGBColor(0xC0, 0x1C, 0x1C), "medium": RGBColor(0xB0, 0x80, 0x00), "low": RGBColor(0x1F, 0x6F, 0xA8)}
_RISK_HEX = {"high": "#C01C1C", "medium": "#B08000", "low": "#1F6FA8"}


def comment_text(f: Finding) -> str:
    lines = [f"[{f.risk.upper()} · {f.score:.1f}]"]
    lines += [f"- {r.rule_id}: {r.message} (+{r.weight:.1f})" for r in f.reasons]
    if f.suggestions:
        lines.append("교정안:")
        lines += [f"{i}) {s.text}" + (f" — {s.note}" if s.note else "") for i, s in enumerate(f.suggestions, 1)]
    return "\n".join(lines)


def write_json(result: AnalysisResult, path: str | Path) -> None:
    Path(path).write_text(json.dumps(result.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")


def _summary_rows(result: AnalysisResult) -> list[tuple[str, str]]:
    s = result.summary()
    rc = s["risk_counts"]
    rows = [
        ("분석 문장 수", str(s["sentences"])),
        ("High / Medium / Low", f'{rc["high"]} / {rc["medium"]} / {rc["low"]}'),
        ("문서 위험 점수 (0–100)", str(s["document_risk_score"])),
    ]
    if result.reference_checks:
        rows.append(("참고문헌 문제", f'{s["reference_issues"]} / {len(result.reference_checks)}'))
    return rows


def write_docx_report(result: AnalysisResult, path: str | Path, min_risk: str = "low") -> None:
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("Humanizer 분석 리포트", 0)
    d.add_paragraph(Path(result.manuscript.path).name)
    d.add_paragraph().add_run(DISCLAIMER).italic = True

    d.add_heading("요약", 1)
    t = d.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in _summary_rows(result):
        cells = t.add_row().cells
        cells[0].text, cells[1].text = k, v

    s = result.summary()
    if s["per_section"]:
        d.add_heading("섹션별 분포", 1)
        t = d.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ("섹션", "High", "Medium", "Low")):
            c.text = h
        for sec, cnt in s["per_section"].items():
            cells = t.add_row().cells
            cells[0].text = sec
            for i, k in enumerate(("high", "medium", "low"), 1):
                cells[i].text = str(cnt[k])
    if s["top_rules"]:
        d.add_heading("자주 걸린 규칙", 1)
        for rule, n in s["top_rules"]:
            d.add_paragraph(f"{rule}: {n}회", style="List Bullet")

    d.add_heading("문장별 상세", 1)
    for n, f in enumerate(result.by_risk(min_risk), 1):
        p = d.add_paragraph()
        head = p.add_run(f"#{n} [{f.risk.upper()} {f.score:.1f}] {f.section} · 문단 {f.para_index}")
        head.bold = True
        head.font.color.rgb = _RISK_RGB.get(f.risk)
        d.add_paragraph(f.text)
        for r in f.reasons:
            d.add_paragraph(f"{r.rule_id}: {r.message} (+{r.weight:.1f})", style="List Bullet")
        for i, sg in enumerate(f.suggestions, 1):
            q = d.add_paragraph(style="List Number")
            q.add_run(sg.text)
            if sg.note:
                q.add_run(f"  — {sg.note}").italic = True

    if result.reference_checks:
        d.add_heading("참고문헌 검증", 1)
        for rc in result.reference_checks:
            if rc.status != "ok":
                d.add_paragraph(f"[{rc.index + 1}] {rc.status}: {rc.detail}\n{rc.text}", style="List Bullet")
    for note in result.notes:
        d.add_paragraph().add_run(note).italic = True
    d.save(str(path))


def _register_font() -> str:
    """Use a CJK-capable font so Korean text renders in the PDF report."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
        return "HYSMyeongJo-Medium"
    except Exception:  # noqa: BLE001 - fall back to a Latin font
        return "Helvetica"


def write_pdf_report(result: AnalysisResult, path: str | Path, min_risk: str = "low") -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _register_font()
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontName=font, fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    h1 = ParagraphStyle("h1", parent=base["Heading1"], fontName=font, fontSize=15)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontName=font, fontSize=12)

    def para(text: str, style=body):
        return Paragraph(escape(text).replace("\n", "<br/>"), style)

    story = [para("Humanizer 분석 리포트", h1), para(Path(result.manuscript.path).name, small),
             para(DISCLAIMER, small), Spacer(1, 4 * mm), para("요약", h2)]
    tbl = Table([[para(k), para(v)] for k, v in _summary_rows(result)], colWidths=[60 * mm, 100 * mm])
    tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey)]))
    story += [tbl, Spacer(1, 4 * mm), para("문장별 상세", h2)]

    rows = [[para("#"), para("위험"), para("문장 / 근거 / 교정안")]]
    for n, f in enumerate(result.by_risk(min_risk), 1):
        detail = f.text + "\n" + "\n".join(f"• {r.rule_id}: {r.message}" for r in f.reasons)
        if f.suggestions:
            detail += "\n" + "\n".join(f"→ {s.text}" + (f" ({s.note})" if s.note else "") for s in f.suggestions)
        risk = Paragraph(f'<font color="{_RISK_HEX.get(f.risk, "#000000")}">{f.risk.upper()}<br/>{f.score:.1f}</font>', body)
        rows.append([para(str(n)), risk, para(detail)])
    tbl = Table(rows, colWidths=[10 * mm, 18 * mm, 142 * mm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
    ]))
    story.append(tbl)
    bad_refs = [rc for rc in result.reference_checks if rc.status != "ok"]
    if bad_refs:
        story += [Spacer(1, 4 * mm), para("참고문헌 검증", h2)]
        story += [para(f"[{rc.index + 1}] {rc.status}: {rc.detail}\n{rc.text}", small) for rc in bad_refs]
    for note in result.notes:
        story.append(para(note, small))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm).build(story)
