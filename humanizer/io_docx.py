"""Read .docx manuscripts and write an annotated copy (highlights + Word comments)."""
from __future__ import annotations

import copy
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.run import Run

from .manuscript import build_manuscript
from .models import AnalysisResult, Manuscript
from .report import comment_text

HIGHLIGHT = {"high": WD_COLOR_INDEX.RED, "medium": WD_COLOR_INDEX.YELLOW, "low": WD_COLOR_INDEX.TURQUOISE}
_PLAIN_RUN_CHILDREN = {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr"), qn("w:noBreakHyphen")}


def _looks_like_heading(p) -> bool:
    style = (p.style.name or "").lower() if p.style is not None else ""
    if style.startswith(("heading", "title", "subtitle")):
        return True
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and len(p.text.split()) <= 12 and all(r.bold for r in runs)


def read_docx(path: str | Path) -> Manuscript:
    d = docx.Document(str(path))
    blocks = ((p.text, {"docx_paragraph": i}, _looks_like_heading(p)) for i, p in enumerate(d.paragraphs))
    return build_manuscript(str(path), "docx", blocks)


def _is_plain(run: Run) -> bool:
    return all(child.tag in _PLAIN_RUN_CHILDREN for child in run._r)


def _split_at(paragraph, offset: int) -> None:
    pos = 0
    for run in paragraph.runs:
        n = len(run.text)
        if pos < offset < pos + n:
            k = offset - pos
            text = run.text
            new_r = copy.deepcopy(run._r)
            run._r.addnext(new_r)
            run.text = text[:k]
            Run(new_r, paragraph).text = text[k:]
            return
        pos += n


def _isolate(paragraph, start: int, end: int) -> list[Run] | None:
    """Split runs so that [start, end) is covered by whole runs; None if unsafe."""
    if "".join(r.text for r in paragraph.runs) != paragraph.text:
        return None  # hyperlinks or fields shift offsets
    pos = 0
    for run in paragraph.runs:
        n = len(run.text)
        if pos < end and pos + n > start and not _is_plain(run):
            return None
        pos += n
    _split_at(paragraph, start)
    _split_at(paragraph, end)
    covered, pos = [], 0
    for run in paragraph.runs:
        n = len(run.text)
        if n and pos >= start and pos + n <= end:
            covered.append(run)
        pos += n
    return covered or None


def annotate_docx(result: AnalysisResult, out_path: str | Path, min_risk: str = "medium",
                  author: str = "Humanizer") -> int:
    """Copy the source .docx with highlighted risky sentences and a comment on each."""
    d = docx.Document(result.manuscript.path)
    paras = d.paragraphs
    by_index = {p.index: p for p in result.manuscript.paragraphs}
    count = 0
    for f in result.by_risk(min_risk):
        src = by_index[f.para_index].source.get("docx_paragraph")
        if src is None:
            continue
        p = paras[src]
        runs = _isolate(p, f.start, f.end)
        if runs:
            for r in runs:
                r.font.highlight_color = HIGHLIGHT.get(f.risk)
        anchor = runs or [r for r in p.runs if r.text] or None
        if anchor:
            text = comment_text(f) if runs else "(문단 전체에 표시) " + comment_text(f)
            d.add_comment(anchor, text=text, author=author, initials="HZ")
            count += 1
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))
    return count
