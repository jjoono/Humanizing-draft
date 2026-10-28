import json
import zipfile

import pymupdf

from humanizer.manuscript import load_manuscript
from humanizer.pipeline import analyze_file, write_outputs


def test_docx_sections_and_references(sample_docx):
    ms = load_manuscript(sample_docx)
    sections = {p.section for p in ms.body()}
    assert {"abstract", "introduction", "methods"} <= sections
    assert ms.references and "10.1063/1.98799" in ms.references[0]


def test_docx_end_to_end(sample_docx, tmp_path):
    result = analyze_file(sample_docx)
    assert result.summary()["risk_counts"]["high"] >= 3
    paths = write_outputs(result, tmp_path / "out")
    names = {p.name for p in paths}
    assert names == {"draft.annotated.docx", "draft.report.json", "draft.report.docx", "draft.report.pdf"}

    annotated = tmp_path / "out" / "draft.annotated.docx"
    with zipfile.ZipFile(annotated) as z:
        comments = z.read("word/comments.xml").decode()
        body = z.read("word/document.xml").decode()
    assert "LEX_PHRASE" in comments
    assert 'w:highlight w:val="red"' in body
    # Splitting runs must not change the document text.
    import docx
    before = [p.text for p in docx.Document(sample_docx).paragraphs]
    after = [p.text for p in docx.Document(annotated).paragraphs]
    assert before == after

    data = json.loads((tmp_path / "out" / "draft.report.json").read_text())
    assert data["summary"]["sentences"] > 0
    assert any(s["source"] == "dictionary" for f in data["findings"] for s in f["suggestions"])


def test_pdf_end_to_end(sample_pdf, tmp_path):
    result = analyze_file(sample_pdf)
    assert {p.section for p in result.manuscript.body()} >= {"introduction", "methods"}
    assert result.summary()["risk_counts"]["high"] >= 2
    write_outputs(result, tmp_path / "out", formats=("annotated",))
    with pymupdf.open(tmp_path / "out" / "draft.annotated.pdf") as doc:
        annots = [(a.type[1], a.info["content"]) for a in doc[0].annots()]
    assert annots and all(kind == "Highlight" for kind, _ in annots)
    assert "LEX_" in annots[0][1]
