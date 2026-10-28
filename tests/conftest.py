import json

import docx
import pymupdf
import pytest

# Written for these tests; deliberately AI-flavoured vs. plain pre-2020-style prose.
AI_PARA = (
    "Perovskite light-emitting diodes have garnered significant attention in the realm of next-generation displays. "
    "Moreover, the intricate interplay between ion migration and charge balance plays a crucial role in device "
    "stability, highlighting the need for comprehensive strategies. This work underscores the pivotal importance of "
    "interface engineering—a key step toward commercialization."
)
HUMAN_PARA = (
    "The devices were fabricated on ITO-coated glass substrates with a sheet resistance of 15 Ω/sq. "
    "A 40 nm PEDOT:PSS layer was spin-coated and annealed at 150 °C for 15 min. "
    "The maximum EQE of 20.3% was obtained at 5 V, which is close to the value reported by Tang et al. [1]. "
    "As shown in Fig. 2a, the EL spectrum peaks at 520 nm."
)


@pytest.fixture
def sample_docx(tmp_path):
    d = docx.Document()
    d.add_heading("Efficient green perovskite LEDs", 0)
    d.add_heading("Abstract", 1)
    d.add_paragraph(AI_PARA)
    d.add_heading("1. Introduction", 1)
    p = d.add_paragraph()
    p.add_run("Perovskite light-emitting diodes have garnered significant attention ")
    p.add_run("in the realm of").bold = True
    p.add_run(" next-generation displays. Moreover, it is worth noting that efficiency is limited.")
    d.add_heading("2. Experimental", 1)
    d.add_paragraph(HUMAN_PARA)
    d.add_heading("References", 1)
    d.add_paragraph("[1] C. W. Tang, S. A. VanSlyke, Appl. Phys. Lett. 51, 913 (1987). doi:10.1063/1.98799")
    path = tmp_path / "draft.docx"
    d.save(path)
    return path


@pytest.fixture
def sample_pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for text, size in (("Introduction", 14), (AI_PARA, 10), ("Experimental", 14), (HUMAN_PARA, 10)):
        rect = pymupdf.Rect(72, y, 520, y + (30 if size == 14 else 120))
        page.insert_textbox(rect, text, fontsize=size)
        y = rect.y1 + 10
    path = tmp_path / "draft.pdf"
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def corpora(tmp_path):
    human = tmp_path / "human.jsonl"
    ai = tmp_path / "ai.jsonl"
    with open(human, "w") as fh:
        for i in range(5):
            fh.write(json.dumps({"id": f"h{i}", "sections": [
                {"section": "methods", "paragraphs": [HUMAN_PARA]},
                {"section": "introduction", "paragraphs": [
                    "Organic light-emitting diodes are used in displays. Their efficiency depends on outcoupling. "
                    "We studied the effect of the cavity length on the emission."]},
            ]}) + "\n")
    with open(ai, "w") as fh:
        for i in range(5):
            fh.write(json.dumps({"id": f"a{i}", "sections": [{"section": "introduction", "paragraphs": [AI_PARA]}]}) + "\n")
    return human, ai
